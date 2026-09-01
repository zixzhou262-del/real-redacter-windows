from __future__ import annotations

from pathlib import Path
from typing import Any

from offline_anonymizer.core.document import Document, DocumentNode
from offline_anonymizer.exporter.base import DocumentExporter


DEFAULT_CJK_FONT = "PingFang SC"
DEFAULT_LATIN_FONT = "Arial"
CODE_FONT = "Menlo"


class DocxExporter(DocumentExporter):
    supported_extensions = {".docx"}

    def export(self, document: Document, path: str | Path) -> None:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise RuntimeError("python-docx is required to export .docx files") from exc

        docx = DocxDocument()
        _set_document_defaults(docx)
        for node in document.blocks:
            _add_node(docx, node)
        docx.save(str(path))


def _set_document_defaults(docx: Any) -> None:
    try:
        normal = docx.styles["Normal"]
        normal.font.name = DEFAULT_LATIN_FONT
        _set_east_asia_font(normal._element, DEFAULT_CJK_FONT)
        for style_name in ["Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
            if style_name in docx.styles:
                style = docx.styles[style_name]
                style.font.name = DEFAULT_LATIN_FONT
                _set_east_asia_font(style._element, DEFAULT_CJK_FONT)
    except Exception:
        return


def _add_node(docx: Any, node: DocumentNode) -> None:
    if node.node_type == "heading":
        try:
            level = max(1, min(int(node.metadata.get("level", 1)), 9))
        except (TypeError, ValueError):
            level = 1
        paragraph = docx.add_heading("", level=level)
        _add_runs(paragraph, node)
    elif node.node_type == "paragraph":
        paragraph = docx.add_paragraph()
        _add_runs(paragraph, node)
    elif node.node_type == "list":
        ordered = node.metadata.get("list_type") == "ordered"
        style = "List Number" if ordered else "List Bullet"
        for child in node.children:
            paragraph = docx.add_paragraph(style=style)
            _add_runs(paragraph, child)
    elif node.node_type == "table":
        _add_table(docx, node)
    elif node.node_type == "code_block":
        paragraph = docx.add_paragraph()
        run = paragraph.add_run(_xml_safe_text(node.content))
        _set_run_font(run, CODE_FONT, CODE_FONT)
    elif node.node_type == "blank_line":
        docx.add_paragraph("")
    else:
        paragraph = docx.add_paragraph()
        _add_runs(paragraph, node)


def _add_table(docx: Any, node: DocumentNode) -> None:
    rows = [row for row in node.children if row.node_type == "table_row"]
    if not rows:
        return
    column_count = max(len(row.children) for row in rows)
    table = docx.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, cell in enumerate(row.children):
            target = table.cell(row_index, column_index)
            target.text = ""
            paragraphs = [
                child for child in cell.children if child.node_type == "paragraph"
            ]
            if paragraphs:
                for index, paragraph_node in enumerate(paragraphs):
                    paragraph = target.paragraphs[0] if index == 0 else target.add_paragraph()
                    _add_runs(paragraph, paragraph_node)
            else:
                target.text = _xml_safe_text(cell.to_plain_text())
                for paragraph in target.paragraphs:
                    for run in paragraph.runs:
                        _set_run_font(run)


def _add_runs(paragraph: Any, node: DocumentNode) -> None:
    text_nodes = node.text_nodes()
    if not text_nodes:
        run = paragraph.add_run(_xml_safe_text(node.content or node.to_plain_text()))
        _set_run_font(run)
        return

    for text in text_nodes:
        run = paragraph.add_run(_xml_safe_text(text.content))
        _set_run_format(run, text.metadata)


def _set_run_font(
    run: Any,
    latin_font: str = DEFAULT_LATIN_FONT,
    east_asia_font: str = DEFAULT_CJK_FONT,
) -> None:
    latin = _valid_font_name(latin_font) or DEFAULT_LATIN_FONT
    east_asia = _valid_font_name(east_asia_font) or DEFAULT_CJK_FONT
    try:
        run.font.name = latin
    except Exception:
        try:
            run.font.name = DEFAULT_LATIN_FONT
        except Exception:
            pass
    _set_east_asia_font(run._element, east_asia)


def _set_east_asia_font(element: Any, font_name: str) -> None:
    try:
        from docx.oxml.ns import qn

        rpr = element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), _valid_font_name(font_name) or DEFAULT_CJK_FONT)
    except Exception:
        return


def _set_run_format(run: Any, metadata: dict[str, Any]) -> None:
    for attribute in ("bold", "italic", "underline"):
        try:
            setattr(run, attribute, bool(metadata.get(attribute)))
        except Exception:
            pass
    _set_run_font(
        run,
        metadata.get("font_name") or DEFAULT_LATIN_FONT,
        metadata.get("font_east_asia") or DEFAULT_CJK_FONT,
    )


def _valid_font_name(value: Any) -> str | None:
    if not isinstance(value, str):
        return None
    cleaned = "".join(character for character in value.strip() if ord(character) >= 32)
    if not cleaned or len(cleaned) > 128:
        return None
    return cleaned


def _xml_safe_text(value: Any) -> str:
    text = value if isinstance(value, str) else str(value or "")
    return "".join(
        character
        if (
            character in "\t\n\r"
            or 0x20 <= ord(character) <= 0xD7FF
            or 0xE000 <= ord(character) <= 0xFFFD
            or 0x10000 <= ord(character) <= 0x10FFFF
        )
        else "�"
        for character in text
    )
