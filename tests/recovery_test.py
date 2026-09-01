import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from offline_anonymizer.anonymizer.engine import AnonymizationEngine
from offline_anonymizer.anonymizer.token_registry import TokenRegistry
from offline_anonymizer.core.document import Document, container_node, text_node
from offline_anonymizer.renderer.html_renderer import HtmlRenderer
from offline_anonymizer.recovery.recovery_engine import (
    HtmlRecoveryEngine,
    MaskedDocumentRecoveryEngine,
)


def token(token_id, label, token_type="company"):
    return (
        f'<span class="masked-token" contenteditable="false" '
        f'data-token-id="{token_id}" data-token-type="{token_type}">{label}</span>'
    )


class RecoveryEngineTest(unittest.TestCase):
    def test_recovers_tokens_by_token_id(self):
        html = token("company_001", "[企业001]") + "收入" + token(
            "money_001", "[金额001]", "money"
        )
        mapping = {
            "company_001": {"type": "company", "original": "北京科技有限公司"},
            "money_001": {"type": "money", "original": "500万元"},
        }

        result = HtmlRecoveryEngine().recover_html(html, mapping)

        self.assertEqual(result.document.to_plain_text(), "北京科技有限公司收入500万元")
        self.assertEqual(result.report.to_dict()["restored_tokens"], 2)
        self.assertEqual(result.report.to_dict()["missing_tokens"], 0)

    def test_recovers_after_user_edits_plain_text(self):
        html = (
            token("company_001", "[企业001]")
            + "收入"
            + token("money_001", "[金额001]", "money")
            + "，新增项目"
        )
        mapping = {
            "company_001": {"type": "company", "original": "北京科技有限公司"},
            "money_001": {"type": "money", "original": "500万元"},
        }

        result = HtmlRecoveryEngine().recover_html(html, mapping)

        self.assertEqual(
            result.document.to_plain_text(),
            "北京科技有限公司收入500万元，新增项目",
        )

    def test_recovers_repeated_token_instances(self):
        html = (
            token("company_001", "[企业001]")
            + "投资"
            + token("company_001", "[企业001]")
            + "并合作"
        )
        mapping = {
            "company_001": {"type": "company", "original": "北京科技有限公司"},
        }

        result = HtmlRecoveryEngine().recover_html(html, mapping)

        self.assertEqual(
            result.document.to_plain_text(),
            "北京科技有限公司投资北京科技有限公司并合作",
        )
        self.assertEqual(result.report.used_token_ids["company_001"], 2)

    def test_recovers_when_token_positions_are_swapped(self):
        html = token("company_002", "[企业002]") + "和" + token(
            "company_001", "[企业001]"
        )
        mapping = {
            "company_001": {"type": "company", "original": "企业A"},
            "company_002": {"type": "company", "original": "企业B"},
        }

        result = HtmlRecoveryEngine().recover_html(html, mapping)

        self.assertEqual(result.document.to_plain_text(), "企业B和企业A")

    def test_recovery_ignores_token_label_text(self):
        html = token("company_001", "用户改坏的显示文本")
        mapping = {
            "company_001": {"type": "company", "original": "北京科技有限公司"},
        }

        result = HtmlRecoveryEngine().recover_html(html, mapping)

        self.assertEqual(result.document.to_plain_text(), "北京科技有限公司")

    def test_reports_missing_and_unused_tokens(self):
        html = token("company_999", "[企业999]")
        mapping = {
            "company_001": {"type": "company", "original": "北京科技有限公司"},
        }

        result = HtmlRecoveryEngine().recover_html(html, mapping)
        report = result.report.to_dict()

        self.assertEqual(result.document.to_plain_text(), "[企业999]")
        self.assertEqual(report["missing_tokens"], 1)
        self.assertEqual(report["missing_token_ids"], ["company_999"])
        self.assertEqual(report["unused_tokens"], 1)
        self.assertEqual(report["unused_token_ids"], ["company_001"])

    def test_recovers_from_html_and_mapping_files(self):
        html = token("company_001", "[企业001]") + "收入稳定"
        mapping = {
            "company_001": {"type": "company", "original": "北京科技有限公司"},
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            html_path = Path(tmpdir) / "edited.html"
            mapping_path = Path(tmpdir) / "mapping.json"
            html_path.write_text(html, encoding="utf-8")
            mapping_path.write_text(
                json.dumps(mapping, ensure_ascii=False),
                encoding="utf-8",
            )

            result = HtmlRecoveryEngine().recover_files(html_path, mapping_path)

        self.assertEqual(result.document.to_plain_text(), "北京科技有限公司收入稳定")

    def test_recovers_from_full_rendered_html_after_plain_text_edit(self):
        source = Document.create(
            "report.txt",
            "txt",
            children=[
                container_node(
                    "paragraph",
                    content="北京科技有限公司收入500万元",
                    children=[text_node("北京科技有限公司收入500万元")],
                )
            ],
        )
        registry = TokenRegistry()
        anonymized = AnonymizationEngine(registry=registry).anonymize(source)
        rendered = HtmlRenderer().render(anonymized, registry)
        edited_html = rendered.html.replace("</p>", "，新增项目</p>", 1)

        result = HtmlRecoveryEngine().recover_html(edited_html, rendered.mapping)

        self.assertEqual(
            result.document.to_plain_text(),
            "北京科技有限公司收入500万元，新增项目",
        )

    def test_recovers_from_edited_masked_txt_file(self):
        mapping = {
            "company_001": {
                "type": "company",
                "original": "北京科技有限公司",
                "masked": "[企业001]",
            },
            "money_001": {
                "type": "money",
                "original": "500万元",
                "masked": "[金额001]",
            },
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            masked_path = root / "edited_masked.txt"
            mapping_path = root / "mapping.json"
            masked_path.write_text(
                "[企业001]收入[金额001]，新增[企业001]项目",
                encoding="utf-8",
            )
            mapping_path.write_text(
                json.dumps(mapping, ensure_ascii=False),
                encoding="utf-8",
            )

            result = MaskedDocumentRecoveryEngine().recover_files(
                masked_path,
                mapping_path,
            )

        self.assertEqual(
            result.document.to_plain_text(),
            "北京科技有限公司收入500万元，新增北京科技有限公司项目",
        )
        self.assertEqual(result.report.restored_tokens, 3)

    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed",
    )
    def test_recovers_from_edited_masked_docx_file(self):
        from docx import Document as DocxDocument

        mapping = {
            "company_001": {
                "type": "company",
                "original": "北京科技有限公司",
                "masked": "[企业001]",
            },
            "money_001": {
                "type": "money",
                "original": "500万元",
                "masked": "[金额001]",
            },
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            masked_path = root / "edited_masked.docx"
            mapping_path = root / "mapping.json"
            docx = DocxDocument()
            docx.add_paragraph("[企业001]收入[金额001]，新增[企业001]项目")
            docx.save(masked_path)
            mapping_path.write_text(
                json.dumps(mapping, ensure_ascii=False),
                encoding="utf-8",
            )

            result = MaskedDocumentRecoveryEngine().recover_files(
                masked_path,
                mapping_path,
            )

        self.assertEqual(
            result.document.to_plain_text(),
            "北京科技有限公司收入500万元，新增北京科技有限公司项目",
        )
        self.assertEqual(result.report.restored_tokens, 3)

    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed",
    )
    def test_recovers_masked_docx_tokens_split_across_runs(self):
        from docx import Document as DocxDocument

        mapping = {
            "company_001": {
                "type": "company",
                "original": "北京科技有限公司",
                "masked": "[企业001]",
            },
            "money_001": {
                "type": "money",
                "original": "500万元",
                "masked": "[金额001]",
            },
        }

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            masked_path = root / "split_runs.docx"
            mapping_path = root / "mapping.json"
            docx = DocxDocument()
            paragraph = docx.add_paragraph()
            paragraph.add_run("[企业")
            paragraph.add_run("001]收入[金")
            paragraph.add_run("额001]，备注已修改")
            docx.save(masked_path)
            mapping_path.write_text(
                json.dumps(mapping, ensure_ascii=False),
                encoding="utf-8",
            )

            result = MaskedDocumentRecoveryEngine().recover_files(
                masked_path,
                mapping_path,
            )

        self.assertEqual(
            result.document.to_plain_text(),
            "北京科技有限公司收入500万元，备注已修改",
        )
        self.assertEqual(result.report.restored_tokens, 2)


if __name__ == "__main__":
    unittest.main()
