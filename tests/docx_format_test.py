import importlib.util
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from offline_anonymizer.security.mapping_crypto import read_mapping
from offline_anonymizer.services import DocumentService


@unittest.skipIf(
    importlib.util.find_spec("docx") is None,
    "python-docx is not installed in this Python environment",
)
class DocxFormatPreservationTest(unittest.TestCase):
    def test_anonymize_and_recover_keep_original_docx_format(self):
        from docx import Document as DocxDocument
        from docx.enum.style import WD_STYLE_TYPE
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "合同.docx"
            output_root = root / "outputs"
            recovered_root = root / "recovered"

            docx = DocxDocument()
            section = docx.sections[0]
            section.top_margin = Inches(0.71)
            section.bottom_margin = Inches(0.83)
            style = docx.styles.add_style("合同正文", WD_STYLE_TYPE.PARAGRAPH)
            style.font.name = "Times New Roman"
            style._element.get_or_add_rPr().get_or_add_rFonts().set(
                qn("w:eastAsia"),
                "Songti SC",
            )
            style.font.size = Pt(13)

            paragraph = docx.add_paragraph(style="合同正文")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(17)
            paragraph.add_run("合同金额：")
            number_run = paragraph.add_run("500")
            number_run.bold = True
            number_run.font.size = Pt(16)
            number_run.font.color.rgb = RGBColor(196, 35, 45)
            unit_run = paragraph.add_run("万元")
            unit_run.italic = True
            paragraph.add_run("，请确认。").underline = True

            table = docx.add_table(rows=1, cols=2)
            table.style = "Light Shading Accent 1"
            table.cell(0, 0).text = "项目"
            table.cell(0, 1).text = "金额500万元"
            header = section.header.paragraphs[0]
            header.add_run("页眉金额500万元").italic = True
            footer = section.footer.paragraphs[0]
            footer.add_run("保密文件").bold = True
            docx.save(source)

            processed = DocumentService().process_file(
                source,
                output_root,
                formats=["docx"],
            )
            self.assertTrue(processed.success, processed.to_dict())
            masked = Path(processed.data["documents"][0])
            mapping = Path(processed.data["mapping"])

            masked_docx = DocxDocument(masked)
            masked_paragraph = masked_docx.paragraphs[0]
            self.assertEqual(masked_paragraph.style.name, "合同正文")
            self.assertEqual(masked_paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(masked_paragraph.paragraph_format.space_after.pt, 17)
            self.assertEqual(masked_paragraph.runs[1].font.size.pt, 16)
            self.assertTrue(masked_paragraph.runs[1].bold)
            self.assertEqual(
                str(masked_paragraph.runs[1].font.color.rgb),
                "C4232D",
            )
            self.assertTrue(masked_paragraph.runs[2].italic)
            self.assertIn("[金额001]", masked_paragraph.text)
            self.assertIn("[金额001]", masked_docx.sections[0].header.paragraphs[0].text)
            self.assertEqual(masked_docx.tables[0].style.name, "Light Shading Accent 1")
            self.assertEqual(
                _package_part(source, "word/styles.xml"),
                _package_part(masked, "word/styles.xml"),
            )
            self.assertEqual(
                _package_part(source, "word/theme/theme1.xml"),
                _package_part(masked, "word/theme/theme1.xml"),
            )
            self.assertEqual(len(read_mapping(mapping)), 1)

            masked_docx.paragraphs[0].add_run(" 用户新增内容").bold = True
            masked_docx.save(masked)
            recovered = DocumentService().recover_document(
                masked,
                mapping,
                recovered_root,
                formats=["docx"],
            )
            self.assertTrue(recovered.success, recovered.to_dict())
            recovered_path = Path(recovered.data["documents"][0])
            restored_docx = DocxDocument(recovered_path)
            restored_paragraph = restored_docx.paragraphs[0]

            self.assertIn("合同金额：500万元，请确认。 用户新增内容", restored_paragraph.text)
            self.assertEqual(restored_paragraph.style.name, "合同正文")
            self.assertEqual(restored_paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            self.assertEqual(restored_paragraph.paragraph_format.space_after.pt, 17)
            self.assertEqual(restored_paragraph.runs[1].text, "500")
            self.assertTrue(restored_paragraph.runs[1].bold)
            self.assertEqual(restored_paragraph.runs[1].font.size.pt, 16)
            self.assertEqual(str(restored_paragraph.runs[1].font.color.rgb), "C4232D")
            self.assertEqual(restored_paragraph.runs[2].text, "万元")
            self.assertTrue(restored_paragraph.runs[2].italic)
            self.assertTrue(restored_paragraph.runs[-1].bold)
            self.assertEqual(restored_docx.tables[0].style.name, "Light Shading Accent 1")
            self.assertIn(
                "页眉金额500万元",
                restored_docx.sections[0].header.paragraphs[0].text,
            )
            self.assertIn(
                "保密文件",
                restored_docx.sections[0].footer.paragraphs[0].text,
            )
            self.assertEqual(
                _package_part(masked, "word/styles.xml"),
                _package_part(recovered_path, "word/styles.xml"),
            )

    def test_deleted_docx_token_is_ignored_and_other_tokens_restore(self):
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            source = root / "金额.docx"
            docx = DocxDocument()
            docx.add_paragraph("第一笔500万元，第二笔800万元。")
            docx.save(source)

            processed = DocumentService().process_file(
                source,
                root / "outputs",
                formats=["docx"],
            )
            self.assertTrue(processed.success, processed.to_dict())
            masked = Path(processed.data["documents"][0])
            edited = DocxDocument(masked)
            edited.paragraphs[0].text = edited.paragraphs[0].text.replace("[金额001]", "")
            edited.save(masked)

            recovered = DocumentService().recover_document(
                masked,
                processed.data["mapping"],
                root / "recovered",
                formats=["docx"],
            )
            self.assertTrue(recovered.success, recovered.to_dict())
            restored = DocxDocument(recovered.data["documents"][0])

        self.assertEqual(restored.paragraphs[0].text, "第一笔，第二笔800万元。")
        self.assertEqual(recovered.data["report"]["restored_tokens"], 1)
        self.assertEqual(recovered.data["report"]["unused_tokens"], 1)


def _package_part(path: Path, name: str) -> bytes:
    with ZipFile(path) as package:
        return package.read(name)


if __name__ == "__main__":
    unittest.main()
