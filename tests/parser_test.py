import importlib.util
import tempfile
import unittest
from pathlib import Path

from offline_anonymizer.parser import ParserFactory


class ParserTest(unittest.TestCase):
    def test_txt_parser_creates_paragraph_node(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "example.txt"
            path.write_text("北京科技有限公司于2025年获得500万元贷款\n", encoding="utf-8")

            document = ParserFactory.get_parser("txt").parse(path)

        paragraphs = document.find_by_type("paragraph")
        self.assertEqual(document.root.node_type, "document")
        self.assertEqual(len(paragraphs), 1)
        self.assertEqual(paragraphs[0].content, "北京科技有限公司于2025年获得500万元贷款")
        self.assertEqual(paragraphs[0].children[0].node_type, "text")
        self.assertEqual(paragraphs[0].children[0].tokens, [])

    def test_markdown_parser_creates_heading_and_paragraph(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "example.md"
            path.write_text("# 项目介绍\n\n北京科技有限公司\n", encoding="utf-8")

            document = ParserFactory.get_parser("md").parse(path)

        headings = document.find_by_type("heading")
        paragraphs = document.find_by_type("paragraph")
        self.assertEqual(headings[0].content, "项目介绍")
        self.assertEqual(headings[0].metadata["level"], 1)
        self.assertEqual(paragraphs[0].content, "北京科技有限公司")

    def test_markdown_parser_creates_list_table_and_code_block(self):
        markdown = "\n".join(
            [
                "## 明细",
                "",
                "- 第一项",
                "- 第二项",
                "",
                "| 名称 | 金额 |",
                "| --- | --- |",
                "| 北京科技有限公司 | 500万元 |",
                "",
                "```python",
                "x = 1",
                "```",
            ]
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "example.md"
            path.write_text(markdown, encoding="utf-8")

            document = ParserFactory.get_parser(path).parse(path)

        self.assertEqual(len(document.find_by_type("list")), 1)
        self.assertEqual(len(document.find_by_type("table")), 1)
        self.assertEqual(document.find_by_type("code_block")[0].content, "x = 1")

    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed in this Python environment",
    )
    def test_docx_parser_creates_heading_paragraph_and_table(self):
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "example.docx"
            docx = DocxDocument()
            docx.add_heading("项目介绍", level=1)
            paragraph = docx.add_paragraph()
            paragraph.add_run("北京科技有限公司").bold = True
            paragraph.add_run("于2025年获得500万元贷款").italic = True
            table = docx.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "名称"
            table.cell(0, 1).text = "金额"
            table.cell(1, 0).text = "北京科技有限公司"
            table.cell(1, 1).text = "500万元"
            docx.save(path)

            document = ParserFactory.get_parser("docx").parse(path)

        self.assertEqual(document.find_by_type("heading")[0].content, "项目介绍")
        self.assertIn("北京科技有限公司", document.find_by_type("paragraph")[0].content)
        self.assertEqual(len(document.find_by_type("table")), 1)
        formatted_text_nodes = [
            node
            for node in document.text_nodes()
            if node.content == "北京科技有限公司"
        ]
        self.assertTrue(any(node.metadata.get("bold") for node in formatted_text_nodes))

    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed in this Python environment",
    )
    def test_docx_parser_accepts_unavailable_and_cjk_fonts(self):
        from docx import Document as DocxDocument
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "fonts.docx"
            docx = DocxDocument()
            paragraph = docx.add_paragraph()
            unavailable = paragraph.add_run("特殊字体内容")
            unavailable.font.name = "Font That Is Not Installed"
            cjk = paragraph.add_run("仿宋内容")
            cjk._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋")
            docx.save(path)

            document = ParserFactory.get_parser(path).parse(path)

        nodes = document.text_nodes()
        self.assertEqual(nodes[0].metadata["font_name"], "Font That Is Not Installed")
        self.assertEqual(nodes[1].metadata["font_east_asia"], "仿宋")


if __name__ == "__main__":
    unittest.main()
