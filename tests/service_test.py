import json
import importlib.util
import tempfile
import unittest
from pathlib import Path

from offline_anonymizer.services import ConfigService, DocumentService, TaskService


class ServiceLayerTest(unittest.TestCase):
    def test_process_file_returns_service_result(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "example.txt"
            output_dir = root / "outputs"
            input_file.write_text("北京科技有限公司收入500万元", encoding="utf-8")

            result = DocumentService().process_file(
                input_file,
                output_dir,
                formats=["txt"],
            )
            masked_text = Path(result.data["documents"][0]).read_text(encoding="utf-8")

            self.assertTrue(result.success)
            self.assertEqual(result.message, "文件处理完成")
            self.assertEqual(result.data["html"], "")
            self.assertEqual(len(result.data["documents"]), 1)
            self.assertEqual(Path(result.data["output_dir"]).parent, output_dir)
            self.assertEqual(Path(result.data["mapping"]).parent, Path(result.data["output_dir"]))
            self.assertEqual(Path(result.data["documents"][0]).parent, Path(result.data["output_dir"]))
            self.assertEqual(Path(result.data["documents"][0]).name, "example已脱敏.txt")
            self.assertEqual(Path(result.data["mapping"]).name, "example还原映射文件.json")
            self.assertEqual(Path(result.data["output_dir"]).name, "example 脱敏文件")
        self.assertIn("北京科技有限公司收入[金额001]", masked_text)
        self.assertNotIn("北京科技有限公司收入500万元", masked_text)

    @unittest.skipIf(importlib.util.find_spec("docx") is None, "python-docx is not installed")
    def test_process_file_creates_single_folder_with_only_docx_and_mapping(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "example.txt"
            output_dir = root / "outputs"
            input_file.write_text("北京科技有限公司收入500万元", encoding="utf-8")

            result = DocumentService().process_file(
                input_file,
                output_dir,
                formats=["docx"],
            )
            run_dir = Path(result.data["output_dir"])
            files = sorted(path.name for path in run_dir.iterdir() if path.is_file())

        self.assertTrue(result.success)
        self.assertEqual(run_dir.parent, output_dir)
        self.assertEqual(run_dir.name, "example 脱敏文件")
        self.assertEqual(files, ["example已脱敏.docx", "example还原映射文件.json"])

    def test_process_file_normalizes_quoted_input_path(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "example.txt"
            output_dir = root / "outputs"
            input_file.write_text("北京科技有限公司收入500万元", encoding="utf-8")

            result = DocumentService().process_file(
                f'  "{input_file}"  ',
                output_dir,
                formats=["txt"],
            )
            output_exists = Path(result.data["documents"][0]).exists()

        self.assertTrue(result.success)
        self.assertTrue(output_exists)

    def test_process_batch_returns_batch_result(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            (input_dir / "a.txt").write_text(
                "北京科技有限公司收入500万元",
                encoding="utf-8",
            )
            (input_dir / "b.md").write_text(
                "# 项目\n\n北京科技有限公司收入500万元",
                encoding="utf-8",
            )

            result = DocumentService().process_batch(
                input_dir,
                output_dir,
                formats=["txt"],
                workers=1,
            )

        self.assertTrue(result.success)
        self.assertEqual(result.data["total"], 2)
        self.assertEqual(result.data["success"], 2)
        self.assertTrue(result.data["report_file"].endswith("batch_report.json"))
        first_output = result.data["files"][0]["output_files"]["txt"]
        self.assertTrue(first_output.endswith("已脱敏.txt"))

    def test_recover_document_restores_and_exports(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "example.txt"
            output_dir = root / "outputs"
            recovered_dir = root / "recovered"
            input_file.write_text("北京科技有限公司收入500万元", encoding="utf-8")
            processed = DocumentService().process_file(
                input_file,
                output_dir,
                formats=["txt"],
            )

            result = DocumentService().recover_document(
                processed.data["documents"][0],
                processed.data["mapping"],
                recovered_dir,
                formats=["txt"],
            )

            recovered_text = Path(result.data["documents"][0]).read_text(encoding="utf-8")

        self.assertTrue(result.success)
        self.assertIn("北京科技有限公司收入500万元", recovered_text)
        self.assertEqual(Path(result.data["output_dir"]).name, "example 恢复文件")

    def test_recover_document_accepts_edited_masked_docx(self):
        try:
            from docx import Document as DocxDocument
        except ImportError:
            self.skipTest("python-docx is not installed")

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "example.txt"
            output_dir = root / "outputs"
            recovered_dir = root / "recovered"
            input_file.write_text("北京科技有限公司收入500万元", encoding="utf-8")
            processed = DocumentService().process_file(
                input_file,
                output_dir,
                formats=["docx"],
            )

            masked_docx = Path(processed.data["documents"][0])
            edited = DocxDocument(masked_docx)
            edited.paragraphs[0].add_run("，新增[金额001]项目")
            edited.save(masked_docx)

            result = DocumentService().recover_document(
                masked_docx,
                processed.data["mapping"],
                recovered_dir,
                formats=["txt"],
            )
            recovered_text = Path(result.data["documents"][0]).read_text(encoding="utf-8")

        self.assertTrue(result.success)
        self.assertIn(
            "北京科技有限公司收入500万元，新增500万元项目",
            recovered_text,
        )

    def test_recover_document_allows_deleted_tokens(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "example.txt"
            output_dir = root / "outputs"
            recovered_dir = root / "recovered"
            input_file.write_text("北京科技有限公司收入500万元", encoding="utf-8")
            processed = DocumentService().process_file(
                input_file,
                output_dir,
                formats=["txt"],
            )
            masked_txt = Path(processed.data["documents"][0])
            masked_txt.write_text("北京科技有限公司收入，新增内容", encoding="utf-8")

            result = DocumentService().recover_document(
                masked_txt,
                processed.data["mapping"],
                recovered_dir,
                formats=["txt"],
            )
            recovered_text = Path(result.data["documents"][0]).read_text(encoding="utf-8")

        self.assertTrue(result.success)
        self.assertIn("北京科技有限公司收入，新增内容", recovered_text)
        self.assertNotIn("500万元", recovered_text)
        self.assertEqual(result.data["report"]["unused_tokens"], 1)
        self.assertEqual(Path(result.data["output_dir"]).parent, recovered_dir)

    def test_export_document_returns_service_result(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "example.txt"
            output_dir = root / "outputs"
            output_file = root / "final.txt"
            input_file.write_text("北京科技有限公司收入500万元", encoding="utf-8")
            processed = DocumentService().process_file(
                input_file,
                output_dir,
                formats=["txt"],
            )
            recovered = DocumentService().recover_document(
                processed.data["documents"][0],
                processed.data["mapping"],
                root / "recovered",
                formats=["txt"],
            )

            result = DocumentService().export_document(
                recovered.data["recovered_document"],
                output_file,
            )

        self.assertTrue(result.success)
        self.assertEqual(result.data["documents"], [str(output_file)])

    def test_missing_file_returns_failure_result_without_throwing(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            result = DocumentService().process_file(
                root / "missing.txt",
                root / "outputs",
                formats=["txt"],
            )
            report_exists = Path(result.data["error_report"]).exists()
            output_folder_exists = any(
                path.is_dir() and "脱敏文件" in path.name for path in root.rglob("*")
            )

        self.assertFalse(result.success)
        self.assertIsNotNone(result.error)
        self.assertIn(result.error["code"], {"FILE_ERROR", "ANONYMIZE_ERROR"})
        self.assertTrue(report_exists)
        self.assertFalse(output_folder_exists)

    def test_failed_anonymization_writes_report_without_output_folder(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "broken.docx"
            output_dir = root / "outputs"
            input_file.write_text("not a valid word document", encoding="utf-8")

            result = DocumentService().process_file(
                input_file,
                output_dir,
                formats=["docx"],
            )
            report_path = Path(result.data["error_report"])
            report = json.loads(report_path.read_text(encoding="utf-8"))

            self.assertFalse(result.success)
            self.assertTrue(report_path.exists())
            self.assertEqual(report["operation"], "脱敏")
            self.assertFalse((output_dir / "broken 脱敏文件").exists())

    def test_all_failed_batch_writes_report_without_batch_output_folder(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_dir = root / "documents"
            output_dir = root / "outputs"
            input_dir.mkdir()
            (input_dir / "broken.docx").write_text("not a docx", encoding="utf-8")

            result = DocumentService().process_batch(
                input_dir,
                output_dir,
                formats=["docx"],
                workers=1,
            )
            report_exists = Path(result.data["error_report"]).exists()
            output_folder_exists = (output_dir / "documents 脱敏文件").exists()

            self.assertFalse(result.success)
            self.assertTrue(report_exists)
            self.assertFalse(output_folder_exists)

    def test_config_service_build_and_save_config(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            service = ConfigService()
            built = service.build_config(
                rules={"company": True, "money": False},
                output_dir=root / "outputs",
                formats=["txt"],
                custom_keywords=["苹果"],
            )
            saved_path = root / "config.json"
            saved = service.save_config(built.data["config"], saved_path)
            loaded = service.load_config(saved_path)

        self.assertTrue(built.success)
        self.assertTrue(saved.success)
        self.assertTrue(loaded.success)
        self.assertFalse(loaded.data["config"]["rules"]["money"])
        self.assertEqual(loaded.data["config"]["output"]["default_formats"], ["txt"])
        self.assertEqual(loaded.data["config"]["custom_keywords"], ["苹果"])

    def test_task_service_manages_gui_task_state(self):
        service = TaskService()
        created = service.create_task("example.txt")
        task_id = created.data["task"]["id"]

        updated = service.update_task(
            task_id,
            status="running",
            progress=50,
            tokens=3,
            message="处理中",
        )
        listed = service.get_tasks()
        cleared = service.clear_tasks()

        self.assertTrue(created.success)
        self.assertEqual(updated.data["task"]["status"], "running")
        self.assertEqual(listed.data["tasks"][0]["tokens"], 3)
        self.assertEqual(cleared.data["tasks"], [])

    def test_progress_callback_receives_status_updates(self):
        events = []
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_file = root / "example.txt"
            input_file.write_text("北京科技有限公司收入500万元", encoding="utf-8")

            result = DocumentService().process_file(
                input_file,
                root / "outputs",
                formats=["txt"],
                progress_callback=events.append,
            )

        self.assertTrue(result.success)
        self.assertEqual(events[0]["status"], "running")
        self.assertEqual(events[-1]["status"], "success")


if __name__ == "__main__":
    unittest.main()
