import importlib.util
import tempfile
import unittest
from pathlib import Path

from offline_anonymizer.anonymizer.engine import AnonymizationEngine
from offline_anonymizer.anonymizer.detector import custom_detectors, default_detectors
from offline_anonymizer.anonymizer.token_registry import TokenRegistry
from offline_anonymizer.core.document import Document, container_node, text_node
from offline_anonymizer.parser import ParserFactory


def document_from_text(text):
    return Document.create(
        "memory.txt",
        "txt",
        children=[
            container_node(
                "paragraph",
                content=text,
                children=[text_node(text)],
            )
        ],
    )


class AnonymizerTest(unittest.TestCase):
    def test_company_number_and_money_are_anonymized(self):
        registry = TokenRegistry()
        engine = AnonymizationEngine(registry=registry)
        document = document_from_text("北京科技有限公司于2025年获得500万元贷款。")

        anonymized = engine.anonymize(document)

        self.assertEqual(
            anonymized.to_plain_text(),
            "[企业001]于[数字001]年获得[金额001]贷款。",
        )
        self.assertEqual(
            document.to_plain_text(),
            "北京科技有限公司于2025年获得500万元贷款。",
        )
        self.assertEqual(
            registry.get_token("company_001").original_value,
            "北京科技有限公司",
        )
        self.assertEqual(registry.get_token("money_001").original_value, "500万元")
        text_node_result = anonymized.text_nodes()[0]
        self.assertEqual([span.token_id for span in text_node_result.tokens], [
            "company_001",
            "number_001",
            "money_001",
        ])

    def test_repeated_company_reuses_same_token(self):
        registry = TokenRegistry()
        engine = AnonymizationEngine(registry=registry)
        document = document_from_text("北京科技有限公司投资北京科技有限公司。")

        anonymized = engine.anonymize(document)

        self.assertEqual(anonymized.to_plain_text(), "[企业001]投资[企业001]。")
        self.assertEqual(len(registry.all_tokens()), 1)
        self.assertEqual(
            len(registry.get_token("company_001").occurrences),
            2,
        )

    def test_phone_and_name_are_anonymized(self):
        registry = TokenRegistry()
        engine = AnonymizationEngine(registry=registry)
        document = document_from_text("联系人张三：\n13812345678")

        anonymized = engine.anonymize(document)

        self.assertEqual(anonymized.to_plain_text(), "联系人[姓名001]：\n[电话001]")
        self.assertEqual(registry.get_token("person_001").original_value, "张三")
        self.assertEqual(registry.get_token("phone_001").original_value, "13812345678")

    def test_money_wins_over_number_on_conflict(self):
        registry = TokenRegistry()
        engine = AnonymizationEngine(registry=registry)
        document = document_from_text("贷款金额为50,000,000元。")

        anonymized = engine.anonymize(document)

        self.assertEqual(anonymized.to_plain_text(), "贷款金额为[金额001]。")
        self.assertIsNone(registry.get_token("number_001"))

    def test_money_and_percent_patterns_are_anonymized(self):
        registry = TokenRegistry()
        engine = AnonymizationEngine(registry=registry)
        document = document_from_text("收入人民币50,000,000.00元，利润率12.5%，增长百分之五。")

        anonymized = engine.anonymize(document)

        self.assertEqual(
            anonymized.to_plain_text(),
            "收入[金额001]，利润率[百分比001]，增长[百分比002]。",
        )
        self.assertEqual(registry.get_token("money_001").original_value, "人民币50,000,000.00元")
        self.assertEqual(registry.get_token("percent_001").original_value, "12.5%")
        self.assertEqual(registry.get_token("percent_002").original_value, "百分之五")

    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed",
    )
    def test_docx_split_runs_are_anonymized_at_paragraph_level(self):
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as tmpdir:
            path = Path(tmpdir) / "split.docx"
            docx = DocxDocument()
            paragraph = docx.add_paragraph()
            paragraph.add_run("北京科技有限公司收入")
            paragraph.add_run("500")
            paragraph.add_run("万元，利润率")
            paragraph.add_run("12.5%")
            docx.save(path)
            document = ParserFactory.get_parser(path).parse(path)

            registry = TokenRegistry()
            anonymized = AnonymizationEngine(registry=registry).anonymize(document)

        self.assertEqual(
            anonymized.to_plain_text(),
            "[企业001]收入[金额001]，利润率[百分比001]",
        )
        self.assertEqual(registry.get_token("money_001").original_value, "500万元")
        self.assertEqual(registry.get_token("percent_001").original_value, "12.5%")

    def test_date_and_email_are_anonymized(self):
        registry = TokenRegistry()
        engine = AnonymizationEngine(registry=registry)
        document = document_from_text("签署日期2025-01-01，邮箱abc@test.com。")

        anonymized = engine.anonymize(document)

        self.assertEqual(anonymized.to_plain_text(), "签署日期[日期001]，邮箱[邮箱001]。")
        self.assertEqual(registry.get_token("date_001").original_value, "2025-01-01")
        self.assertEqual(registry.get_token("email_001").original_value, "abc@test.com")

    def test_custom_keyword_is_anonymized_and_reused(self):
        registry = TokenRegistry()
        engine = AnonymizationEngine(
            detectors=[
                *default_detectors(),
                *custom_detectors(keywords=["苹果"]),
            ],
            registry=registry,
        )
        document = document_from_text("苹果公司发布苹果产品。")

        anonymized = engine.anonymize(document)

        self.assertEqual(anonymized.to_plain_text(), "[敏感001]公司发布[敏感001]产品。")
        self.assertEqual(registry.get_token("custom_001").original_value, "苹果")
        self.assertEqual(len(registry.get_token("custom_001").occurrences), 2)


if __name__ == "__main__":
    unittest.main()
