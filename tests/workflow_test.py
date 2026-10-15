import importlib.util
import json
import tempfile
import unittest
from pathlib import Path

from offline_anonymizer.config import load_config
from offline_anonymizer.parser import ParserFactory
from offline_anonymizer.security.mapping_crypto import read_mapping
from offline_anonymizer.workflow import DocumentPipeline


class WorkflowTest(unittest.TestCase):
    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed in this Python environment",
    )
    def test_pipeline_runs_full_single_docx_workflow(self):
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_path = root / "example.docx"
            output_dir = root / "outputs"
            log_dir = root / "logs"
            config_path = root / "config.json"

            docx = DocxDocument()
            docx.add_heading("项目介绍", level=1)
            docx.add_paragraph("北京科技有限公司收入500万元")
            docx.save(input_path)

            config_path.write_text(
                json.dumps(
                    {
                        "paths": {
                            "output_dir": str(output_dir),
                            "log_dir": str(log_dir),
                        },
                        "output": {"default_formats": ["docx"]},
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result = DocumentPipeline(config_path=config_path).process(input_path)

            output_docx = Path(result.output_files["docx"])
            reparsed = ParserFactory.get_parser(output_docx).parse(output_docx)

            self.assertTrue(Path(result.masked_html).exists())
            self.assertTrue(Path(result.mapping_file).exists())
            self.assertTrue(Path(result.masked_document).exists())
            self.assertTrue(Path(result.recovered_document).exists())
            self.assertTrue(output_docx.exists())
            self.assertIn("北京科技有限公司收入[金额001]", reparsed.to_plain_text())
            self.assertNotIn("北京科技有限公司收入500万元", reparsed.to_plain_text())
            self.assertEqual(result.report["status"], "success")
            self.assertEqual(result.report["tokens"], 1)
            self.assertTrue((log_dir / "workflow.jsonl").exists())

    def test_anonymize_respects_rule_config(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_path = root / "example.txt"
            output_dir = root / "outputs"
            log_dir = root / "logs"
            input_path.write_text(
                "北京科技有限公司收入500万元",
                encoding="utf-8",
            )
            config = {
                "paths": {
                    "output_dir": str(output_dir),
                    "log_dir": str(log_dir),
                },
                "rules": {
                    "company": True,
                    "date": True,
                    "email": True,
                    "money": True,
                    "number": True,
                    "percent": True,
                },
            }

            result = DocumentPipeline(config=config).anonymize(input_path)
            html = Path(result.masked_html).read_text(encoding="utf-8")
            mapping = read_mapping(result.mapping_file)

        self.assertIn("北京科技有限公司收入", html)
        self.assertNotIn("company_001", mapping)
        self.assertIn("money_001", mapping)

    def test_workflow_only_uses_v110_supported_rules_even_if_legacy_config_enables_more(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_path = root / "example.txt"
            input_path.write_text(
                "北京科技有限公司于2025-01-01联系abc@test.com，金额500万元，占比12.5%。",
                encoding="utf-8",
            )
            config = {
                "paths": {
                    "output_dir": str(root / "outputs"),
                    "log_dir": str(root / "logs"),
                },
                "rules": {
                    "company": True,
                    "date": True,
                    "email": True,
                    "person": True,
                    "phone": True,
                    "money": True,
                    "number": True,
                    "percent": True,
                    "custom": True,
                },
            }

            result = DocumentPipeline(config=config).anonymize(input_path)
            html = Path(result.masked_html).read_text(encoding="utf-8")
            mapping = read_mapping(result.mapping_file)

        self.assertIn("北京科技有限公司", html)
        self.assertIn("abc@test.com", html)
        self.assertNotIn("company_001", mapping)
        self.assertNotIn("date_001", mapping)
        self.assertNotIn("email_001", mapping)
        self.assertIn("money_001", mapping)
        self.assertIn("percent_001", mapping)

    def test_anonymize_uses_custom_keywords_from_config(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_path = root / "example.txt"
            output_dir = root / "outputs"
            log_dir = root / "logs"
            input_path.write_text(
                "苹果公司采购苹果设备",
                encoding="utf-8",
            )
            config = {
                "paths": {
                    "output_dir": str(output_dir),
                    "log_dir": str(log_dir),
                },
                "custom_keywords": ["苹果"],
            }

            result = DocumentPipeline(config=config).anonymize(input_path)
            html = Path(result.masked_html).read_text(encoding="utf-8")
            mapping = read_mapping(result.mapping_file)

        self.assertEqual(html.count('data-token-id="custom_001"'), 2)
        self.assertIn("[敏感001]", html)
        self.assertEqual(mapping["custom_001"]["original"], "苹果")

    def test_explicit_custom_only_rules_disable_numeric_detectors(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_path = root / "example.txt"
            input_path.write_text(
                "苹果项目金额500万元，占比12.5%，编号2025。",
                encoding="utf-8",
            )
            config = {
                "paths": {
                    "output_dir": str(root / "outputs"),
                    "log_dir": str(root / "logs"),
                },
                "rules": {
                    "money": False,
                    "percent": False,
                    "number": False,
                    "custom": True,
                },
                "custom_keywords": ["苹果"],
            }

            result = DocumentPipeline(config=config).anonymize(input_path)
            html = Path(result.masked_html).read_text(encoding="utf-8")
            mapping = read_mapping(result.mapping_file)

        self.assertIn("[敏感001]", html)
        self.assertIn("项目金额500万元，占比12.5%，编号2025", html)
        self.assertNotIn("money_001", mapping)
        self.assertNotIn("percent_001", mapping)
        self.assertNotIn("number_001", mapping)
        self.assertEqual(mapping["custom_001"]["original"], "苹果")

    def test_load_config_merges_user_config_with_defaults(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            config_path = Path(tmpdir) / "config.json"
            config_path.write_text(
                json.dumps({"rules": {"company": False}}),
                encoding="utf-8",
            )

            config = load_config(config_path)

        self.assertFalse(config["rules"]["company"])
        self.assertTrue(config["rules"]["money"])
        self.assertTrue(config["rules"]["percent"])
        self.assertTrue(config["rules"]["custom"])
        self.assertIn("default_formats", config["output"])


if __name__ == "__main__":
    unittest.main()
