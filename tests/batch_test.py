import importlib.util
import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from offline_anonymizer.batch import BatchProcessor
from offline_anonymizer.batch import report
from offline_anonymizer.parser import ParserFactory
from offline_anonymizer.services import DocumentService


class BatchProcessorTest(unittest.TestCase):
    def test_memory_peak_is_optional_when_platform_module_is_unavailable(self):
        with patch.object(report, "resource", None):
            self.assertIsNone(report._memory_peak_kb())

    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed in this Python environment",
    )
    def test_batch_processes_supported_files_and_preserves_subdirectories(self):
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_dir = root / "input"
            output_dir = root / "output"
            nested = input_dir / "project"
            nested.mkdir(parents=True)

            (input_dir / "a.txt").write_text(
                "北京科技有限公司收入500万元",
                encoding="utf-8",
            )
            (input_dir / "b.md").write_text(
                "# 项目介绍\n\n北京科技有限公司收入500万元",
                encoding="utf-8",
            )
            docx = DocxDocument()
            docx.add_heading("项目介绍", level=1)
            docx.add_paragraph("北京科技有限公司收入500万元")
            docx.save(nested / "c.docx")
            (input_dir / "ignore.pdf").write_text("ignored", encoding="utf-8")

            result = BatchProcessor(max_workers=1).process_directory(
                input_dir,
                output_dir,
                output_formats=["txt"],
            )

            report_path = output_dir / "reports" / "batch_report.json"
            report = json.loads(report_path.read_text(encoding="utf-8"))

            self.assertEqual(result.total, 3)
            self.assertEqual(result.success, 3)
            self.assertEqual(result.failed, 0)
            self.assertEqual(result.ignored, 1)
            self.assertTrue((output_dir / "a已脱敏.txt").exists())
            self.assertTrue((output_dir / "a还原映射文件.json").exists())
            self.assertTrue((output_dir / "b已脱敏.txt").exists())
            self.assertTrue((output_dir / "b还原映射文件.json").exists())
            self.assertTrue((output_dir / "project" / "c已脱敏.txt").exists())
            self.assertTrue((output_dir / "project" / "c还原映射文件.json").exists())
            masked_text = (output_dir / "project" / "c已脱敏.txt").read_text(
                encoding="utf-8"
            )
            self.assertIn("北京科技有限公司收入[金额001]", masked_text)
            self.assertTrue((output_dir / "批量还原映射文件.json").exists())
            self.assertEqual(report["batch_mapping_file"], str(output_dir / "批量还原映射文件.json"))
            self.assertEqual(report["success"], 3)
            self.assertEqual(report["ignored_files"], ["ignore.pdf"])

    def test_batch_outputs_can_be_merged_and_recovered_with_batch_mapping(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_dir = root / "input"
            output_dir = root / "output"
            recovered_dir = root / "recovered"
            input_dir.mkdir()
            originals = [
                "第一份文档：项目金额100万元，占比10%。",
                "第二份文档：项目金额200万元，占比20%。",
                "第三份文档：项目金额300万元，占比30%。",
                "第四份文档：项目金额400万元，占比40%。",
            ]
            for index, text in enumerate(originals, start=1):
                (input_dir / f"doc{index}.txt").write_text(text, encoding="utf-8")

            result = DocumentService().process_batch(
                input_dir,
                output_dir,
                formats=["txt"],
                workers=1,
            )
            batch_output_dir = Path(result.data["output_dir"])
            masked_files = sorted(batch_output_dir.glob("doc*已脱敏.txt"))
            merged = root / "A.txt"
            merged.write_text(
                "\n".join(path.read_text(encoding="utf-8") for path in masked_files),
                encoding="utf-8",
            )
            recovered = DocumentService().recover_document(
                merged,
                result.data["batch_mapping_file"],
                recovered_dir,
                formats=["txt"],
            )
            recovered_text = Path(recovered.data["documents"][0]).read_text(
                encoding="utf-8"
            )

        self.assertTrue(result.success)
        self.assertTrue(recovered.success)
        for text in originals:
            self.assertIn(text, recovered_text)

    def test_batch_continues_when_one_file_fails(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            (input_dir / "good.txt").write_text(
                "北京科技有限公司收入500万元",
                encoding="utf-8",
            )
            (input_dir / "bad.docx").write_text(
                "not a real docx",
                encoding="utf-8",
            )

            result = BatchProcessor(max_workers=1).process_directory(
                input_dir,
                output_dir,
                output_formats=["txt"],
            )
            report = json.loads(
                (output_dir / "reports" / "batch_report.json").read_text(
                    encoding="utf-8"
                )
            )

        self.assertEqual(result.total, 2)
        self.assertEqual(result.success, 1)
        self.assertEqual(result.failed, 1)
        statuses = {Path(item["file"]).name: item["status"] for item in report["files"]}
        self.assertEqual(statuses["good.txt"], "success")
        self.assertEqual(statuses["bad.docx"], "failed")

    def test_batch_supports_limited_parallel_workers(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            for index in range(10):
                (input_dir / f"file_{index}.txt").write_text(
                    f"北京科技有限公司收入{index}万元",
                    encoding="utf-8",
                )

            started = time.perf_counter()
            result = BatchProcessor(max_workers=4).process_directory(
                input_dir,
                output_dir,
                output_formats=["txt"],
            )
            duration = time.perf_counter() - started

        self.assertEqual(result.total, 10)
        self.assertEqual(result.success, 10)
        self.assertEqual(result.failed, 0)
        self.assertGreaterEqual(result.duration_seconds, 0)
        self.assertGreaterEqual(duration, 0)

    def test_batch_can_process_100_simulated_txt_files(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            input_dir = root / "input"
            output_dir = root / "output"
            input_dir.mkdir()
            for index in range(100):
                (input_dir / f"file_{index}.txt").write_text(
                    f"北京科技有限公司收入{index}万元",
                    encoding="utf-8",
                )

            result = BatchProcessor(max_workers=4).process_directory(
                input_dir,
                output_dir,
                output_formats=["txt"],
            )

        self.assertEqual(result.total, 100)
        self.assertEqual(result.success, 100)
        self.assertEqual(result.failed, 0)
        self.assertGreater(result.average_seconds_per_file, 0)


if __name__ == "__main__":
    unittest.main()
