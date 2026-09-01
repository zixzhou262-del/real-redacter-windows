import importlib.util
import tempfile
import unittest
from pathlib import Path

from offline_anonymizer.anonymizer.engine import AnonymizationEngine
from offline_anonymizer.anonymizer.token_registry import TokenRegistry
from offline_anonymizer.core.document import Document, container_node, text_node
from offline_anonymizer.exporter import DocxExporter, ExporterFactory, MarkdownExporter, TxtExporter
from offline_anonymizer.parser import ParserFactory
from offline_anonymizer.recovery import HtmlRecoveryEngine
from offline_anonymizer.renderer import HtmlRenderer


def sample_document():
    return Document.create(
        "sample.txt",
        "txt",
        children=[
            container_node(
                "heading",
                content="项目介绍",
                children=[text_node("项目介绍")],
                metadata={"level": 1},
            ),
            container_node(
                "paragraph",
                content="北京科技有限公司收入500万元",
                children=[text_node("北京科技有限公司收入500万元")],
            ),
        ],
    )


class ExporterTest(unittest.TestCase):
    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed in this Python environment",
    )
    def test_docx_exporter_writes_file(self):
        from docx import Document as DocxDocument
        from docx.oxml.ns import qn

        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "test.docx"

            DocxExporter().export(sample_document(), output)
            exported = DocxDocument(output)
            first_run = exported.paragraphs[0].runs[0]
            east_asia_font = first_run._element.rPr.rFonts.get(qn("w:eastAsia"))

            self.assertTrue(output.exists())
            self.assertGreater(output.stat().st_size, 0)
            self.assertEqual(east_asia_font, "PingFang SC")

    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed in this Python environment",
    )
    def test_docx_exporter_falls_back_for_invalid_font_and_xml_characters(self):
        from docx import Document as DocxDocument
        from docx.oxml.ns import qn

        document = Document.create(
            "font.docx",
            "docx",
            children=[
                container_node(
                    "paragraph",
                    children=[
                        text_node(
                            "内容\x00继续",
                            metadata={
                                "font_name": "Invalid\x00Font",
                                "font_east_asia": "\x00",
                            },
                        )
                    ],
                )
            ],
        )
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "font.docx"
            DocxExporter().export(document, output)
            exported = DocxDocument(output)
            run = exported.paragraphs[0].runs[0]
            east_asia = run._element.rPr.rFonts.get(qn("w:eastAsia"))
            text = exported.paragraphs[0].text

        self.assertEqual(run.font.name, "InvalidFont")
        self.assertEqual(east_asia, "PingFang SC")
        self.assertEqual(text, "内容�继续")

    def test_markdown_exporter_writes_heading_and_paragraph(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "test.md"

            MarkdownExporter().export(sample_document(), output)

            content = output.read_text(encoding="utf-8")
            self.assertIn("# 项目介绍", content)
            self.assertIn("北京科技有限公司收入500万元", content)

    def test_txt_exporter_writes_plain_text(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            output = Path(tmpdir) / "test.txt"

            TxtExporter().export(sample_document(), output)

            self.assertEqual(
                output.read_text(encoding="utf-8"),
                "项目介绍\n\n北京科技有限公司收入500万元",
            )

    def test_markdown_exporter_writes_list_table_and_code_block(self):
        document = Document.create(
            "sample.md",
            "md",
            children=[
                container_node(
                    "list",
                    children=[
                        container_node(
                            "list_item",
                            content="第一项",
                            children=[text_node("第一项")],
                        )
                    ],
                    metadata={"list_type": "bullet"},
                ),
                container_node(
                    "table",
                    children=[
                        container_node(
                            "table_row",
                            children=[
                                container_node("table_cell", content="名称"),
                                container_node("table_cell", content="金额"),
                            ],
                        ),
                        container_node(
                            "table_row",
                            children=[
                                container_node("table_cell", content="北京科技有限公司"),
                                container_node("table_cell", content="500万元"),
                            ],
                        ),
                    ],
                ),
                container_node(
                    "code_block",
                    content="x = 1",
                    children=[text_node("x = 1")],
                    metadata={"language": "python"},
                ),
            ],
        )

        markdown = MarkdownExporter().supports(".md")
        with tempfile.TemporaryDirectory() as tmpdir:
            rendered_path = Path(tmpdir) / "test.md"
            MarkdownExporter().export(document, rendered_path)
            content = rendered_path.read_text(encoding="utf-8")

        self.assertTrue(markdown)
        self.assertIn("- 第一项", content)
        self.assertIn("| 名称 | 金额 |", content)
        self.assertIn("```python\nx = 1\n```", content)

    @unittest.skipIf(
        importlib.util.find_spec("docx") is None,
        "python-docx is not installed in this Python environment",
    )
    def test_end_to_end_docx_parse_anonymize_html_recover_export(self):
        from docx import Document as DocxDocument

        with tempfile.TemporaryDirectory() as tmpdir:
            source_path = Path(tmpdir) / "example.docx"
            output_path = Path(tmpdir) / "output.docx"

            source_docx = DocxDocument()
            source_docx.add_heading("项目介绍", level=1)
            source_docx.add_paragraph("北京科技有限公司收入500万元")
            source_docx.save(source_path)

            parsed = ParserFactory.get_parser(source_path).parse(source_path)
            registry = TokenRegistry()
            anonymized = AnonymizationEngine(registry=registry).anonymize(parsed)
            rendered = HtmlRenderer().render(anonymized, registry)
            recovered = HtmlRecoveryEngine().recover_html(
                rendered.html,
                rendered.mapping,
            ).document

            ExporterFactory.get_exporter(output_path).export(recovered, output_path)
            self.assertTrue(output_path.exists())

            reparsed = ParserFactory.get_parser(output_path).parse(output_path)

        self.assertIn("项目介绍", reparsed.to_plain_text())
        self.assertIn("北京科技有限公司收入500万元", reparsed.to_plain_text())


if __name__ == "__main__":
    unittest.main()
